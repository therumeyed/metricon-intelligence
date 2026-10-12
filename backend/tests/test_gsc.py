import openpyxl
from sqlalchemy import func, select

from app.models import GscImport, GscPageRow, GscQueryRow
from app.services import gsc


def make_workbook(path, queries=None, pages=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Query", "Clicks", "Impressions", "CTR", "Position"])
    for r in queries or [("house and land packages melbourne", 10, 200, 0.05, 6.5), ("display homes near me", 5, 100, 0, 8), ("curtain cleaning vaucluse", 0, 9, 0, 40.0),
                         ("knock down rebuild cost", 1, 50, 0.02, 12.2), ("lumiere", 3, 60, 0.05, 3.0)]:
        ws.append(list(r))
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["Page", "Clicks", "Impressions", "CTR", "Position"])
    for r in pages or [("https://www.metricon.com.au/house-and-land/packages/melbourne", 7, 700, 0.01, 11.8), ("https://www.metricon.com.au/house-designs/melbourne/lumiere", 2, 90, 0.02, 9.0),
                       ("https://www.metricon.com.au/latest-offers/southern-nsw", 1, 20, 0.05, 4.0)]:
        ws2.append(list(r))
    wb.save(path)


def test_import_both_tabs_typed_and_separate(seeded, tmp_path):
    p = tmp_path / "gsc.xlsx"
    make_workbook(p)
    r = gsc.import_workbook(seeded, p, period_note="reported YTD to 10 Oct; not a column")
    seeded.commit()
    assert r["created"] and r["query_rows"] == 5 and r["page_rows"] == 3
    q = seeded.scalar(select(GscQueryRow).where(GscQueryRow.query == "house and land packages melbourne"))
    assert isinstance(q.clicks, int) and isinstance(q.ctr, float) and q.position == 6.5
    # integer-looking CTR in the source ('0') is preserved as a number, not dropped
    assert seeded.scalar(select(GscQueryRow.ctr).where(GscQueryRow.query == "display homes near me")) == 0
    # the two datasets are separate tables with no join key between them
    assert "page" not in GscQueryRow.__table__.columns and "query" not in GscPageRow.__table__.columns
    imp = seeded.scalar(select(GscImport))
    assert imp.reported_period_start is None and "not a column" in imp.period_note


def test_import_is_idempotent_even_if_file_bytes_differ(seeded, tmp_path):
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    make_workbook(a)
    make_workbook(b)  # same data, different file bytes (timestamps)
    r1 = gsc.import_workbook(seeded, a)
    seeded.commit()
    r2 = gsc.import_workbook(seeded, b)
    seeded.commit()
    assert r1["created"] and not r2["created"] and r1["import_id"] == r2["import_id"]
    assert seeded.scalar(select(func.count()).select_from(GscQueryRow)) == 5


def test_new_export_does_not_overwrite_history(seeded, tmp_path):
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    make_workbook(a)
    make_workbook(b, queries=[("house and land packages melbourne", 99, 999, 0.1, 5.0)])
    gsc.import_workbook(seeded, a)
    gsc.import_workbook(seeded, b)
    seeded.commit()
    assert seeded.scalar(select(func.count()).select_from(GscImport)) == 2
    assert seeded.scalar(select(func.count()).select_from(GscQueryRow)) == 6  # first import's rows are untouched


def test_wrong_headers_rejected(seeded, tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["Keyword", "Clicks"])
    wb.create_sheet("Sheet2").append(["Page", "Clicks", "Impressions", "CTR", "Position"])
    p = tmp_path / "bad.xlsx"
    wb.save(p)
    import pytest

    with pytest.raises(ValueError, match="expected columns"):
        gsc.import_workbook(seeded, p)


def test_theme_mapping_exclusions_and_named_designs(seeded, tmp_path):
    p = tmp_path / "g.xlsx"
    make_workbook(p)
    r = gsc.import_workbook(seeded, p)
    seeded.commit()
    rows = {x.query: x for x in seeded.scalars(select(GscQueryRow).where(GscQueryRow.import_id == r["import_id"]))}
    from app.models import Theme

    slug = lambda row: seeded.scalar(select(Theme.slug).where(Theme.id == row.theme_id)) if row.theme_id else None
    assert slug(rows["house and land packages melbourne"]) == "house-and-land"
    assert slug(rows["knock down rebuild cost"]) == "knockdown-rebuild"  # specific theme beats 'costs'
    assert rows["curtain cleaning vaucluse"].excluded  # off-topic, not mapped
    assert slug(rows["lumiere"]) == "designs-floorplans"  # named design learned from the page tab
    assert rows["house and land packages melbourne"].state_hint == "VIC"


def test_representative_terms_are_few_and_generic(seeded, tmp_path):
    p = tmp_path / "g.xlsx"
    qs = [(f"house and land packages {c}", 1, 100 + i, 0.01, 5.0) for i, c in enumerate(["melbourne", "sydney", "adelaide"])] + [("house and land packages", 9, 900, 0.01, 6.0)]
    make_workbook(p, queries=qs)
    r = gsc.import_workbook(seeded, p)
    gsc.propose_representative_terms(seeded, r["import_id"])
    from app.models import ThemeKeyword

    kws = {k.keyword: k for k in seeded.scalars(select(ThemeKeyword))}
    assert kws["house and land packages"].gsc_impressions == 900  # existing seed term gets its GSC weight, no duplicate row
    assert not any(c in k for k in kws for c in ("melbourne", "sydney", "adelaide"))  # place-specific terms are never proposed for state/national Trends


def test_remap_keeps_raw_values(seeded, tmp_path):
    p = tmp_path / "g.xlsx"
    make_workbook(p)
    r = gsc.import_workbook(seeded, p)
    before = [(x.query, x.clicks, x.impressions, x.ctr, x.position) for x in seeded.scalars(select(GscQueryRow).order_by(GscQueryRow.id))]
    gsc.remap_import(seeded, r["import_id"])
    seeded.commit()
    assert before == [(x.query, x.clicks, x.impressions, x.ctr, x.position) for x in seeded.scalars(select(GscQueryRow).order_by(GscQueryRow.id))]
